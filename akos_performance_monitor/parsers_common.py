"""Shared spreadsheet helpers and driver-table discovery."""
from __future__ import annotations

import csv
import re
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None

SUPPORTED_TABULAR = {".csv", ".xlsx", ".xls"}

DRIVER_ALIASES = (
    "Driver",
    "Driver Name",
    "Delivery Associate",
    "DA Name",
    "Associate Name",
    "Transporter",
    "Transporter Name",
    "Name",
)


def norm(value) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())


def value_float(value):
    if value is None or str(value).strip() in ("", "-", "N/A", "NA", "null"):
        return None
    s = str(value).strip().replace(",", "")
    if s.endswith("%"):
        return float(s[:-1])
    return float(s)


def read_tables(path: Path):
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            yield path.name, list(csv.reader(f))
    elif suffix in {".xlsx", ".xls"}:
        if load_workbook is None:
            raise RuntimeError("Install openpyxl: pip install openpyxl")
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in wb:
                yield sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            wb.close()
    else:
        return


def find_header_table(path: Path, required_aliases: list[str] | tuple[str, ...], scan_rows: int = 80):
    """Return (sheet, header_norms, records) for first table matching any required alias."""
    wanted = {norm(x) for x in required_aliases}
    for sheet, rows in read_tables(path):
        for i, cells in enumerate(rows[:scan_rows]):
            headers = [norm(x) for x in cells]
            if not any(h in wanted for h in headers):
                continue
            records = []
            for row in rows[i + 1 :]:
                if not any(str(c or "").strip() for c in row):
                    continue
                record = {}
                for k in range(min(len(headers), len(row))):
                    if headers[k]:
                        record[headers[k]] = row[k]
                # Skip aggregate rows
                name_keys = {norm(a) for a in DRIVER_ALIASES}
                label = ""
                for nk in name_keys:
                    if nk in record:
                        label = str(record[nk] or "").strip()
                        break
                if label and norm(label) in {"total", "average", "subtotal", "grandtotal", "dsp"}:
                    continue
                records.append(record)
            if records:
                return sheet, headers, records
    return None, [], []


def get_cell(record: dict, aliases: list[str] | tuple[str, ...]):
    for alias in aliases:
        key = norm(alias)
        if key in record:
            return record[key]
    return None


def find_driver_table(path: Path, aliases: dict | None = None):
    driver_aliases = (aliases or {}).get("driver") if aliases else None
    driver_aliases = list(driver_aliases) if driver_aliases else list(DRIVER_ALIASES)
    sheet, _headers, records = find_header_table(path, driver_aliases)
    return sheet, records
